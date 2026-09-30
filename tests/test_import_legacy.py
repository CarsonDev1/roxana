import os
from pathlib import Path

import pytest
from openpyxl import Workbook

import import_legacy
from common import read_records

REAL_CANDIDATES = [os.environ.get("ROXANA_LEGACY_XLSX", ""),
                   r"C:\Users\buitr\Downloads\Roxana_Plaza_Tong_hop_vu_viec (1) (1).xlsx",
                   r"C:\Users\tinhbt\Downloads\Roxana_Plaza_Tong_hop_vu_viec (1).xlsx"]
REAL = next((Path(p) for p in REAL_CANDIDATES if p and Path(p).is_file()), Path(REAL_CANDIDATES[1]))


def make_legacy(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tóm tắt vụ việc"
    ws["A1"] = "TỔNG HỢP VỤ VIỆC DỰ ÁN ROXANA PLAZA"
    for i, (a, b, c) in enumerate([
        ("2017 - 2020", "Naviland bán 1.082 căn hộ Roxana Plaza.", "CafeF / NLD — Kết luận thanh tra"),
        ("22/1/2021", "Bà Phạm Thị Ngọc Liên được bầu làm Tổng Giám đốc Naviland.", "Doanh nghiệp Hội nhập — Bài 3"),
        ("Công an Bình Dương", "Công an tỉnh Bình Dương tiếp nhận tố giác.", "Kinh tế Đô thị"),
        ("22/8/2026", "Buổi đối thoại.", 'Ảnh "Giấy mời" đăng trong nhóm cư dân trên Facebook'),
    ], 5):
        ws.cell(i, 1, a), ws.cell(i, 2, b), ws.cell(i, 3, c)
    ws["A10"] = "LƯU Ý QUAN TRỌNG: Đây là tranh chấp dân sự/hành chính đang trong quá trình xử lý."

    ws = wb.create_sheet("Các bên liên quan")
    for i, row in enumerate([
        ("Công ty TNHH XD-DV-TM-Đầu tư BĐS Tường Phong", "Doanh nghiệp", "Chủ đầu tư chính thức.", "CafeF"),
        ("Bà Phạm Thị Ngọc Liên", "Cá nhân", "Vợ ông Lầu Nam Tường.", "Doanh nghiệp Hội nhập"),
        ("Ông Hoàng Tùng", "Cá nhân", "Được báo Dân Việt nêu tên.", "Dân Việt"),
    ], 5):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)

    ws = wb.create_sheet("Bài viết MXH nổi bật")
    for i, row in enumerate([
        (1, "Lê Trọng Chiến", "https://www.facebook.com/100000273215197",
         "ROXANA PLAZA - HÀNH TRÌNH ĐÒI NHÀ (2.425 thành viên)", "~1/6/2026",
         "Những ai chưa biết vào đâu để đồng hành đòi nhà Roxana thì tham gia nhóm này.", "10 thích",
         "Thông tin/kêu gọi hành động", "https://www.facebook.com/groups/427692059062534/",
         "Không lấy được permalink bài riêng lẻ"),
        (2, "Huynh Bich Diem", "https://www.facebook.com/100009329476350", "CỘNG ĐỒNG CƯ DÂN ROXANA PLAZA",
         "2/12/2024", "Nơi hội tụ tinh hoa???? Xạo, trả nhà cho người dân Roxana.", "4 thích, 4 bl",
         "GAY GẮT / chửi bới", "https://www.facebook.com/photo/?fbid=3966520550335555&set=pcb.1115414933476660", None),
        (3, "Ls. Trần Minh Cường", "(không lấy được)", "REVIEW BẤT ĐỘNG SẢN", "29/7/2026, 15:38",
         "MUA NHÀ 6 NĂM... Công ty CP N. bị kiện, căn hộ Roxana.", "1 thích",
         "Thông tin pháp lý (luật sư, kèm văn bản toà án)", "(không có)", "Chưa lấy được link"),
        (4, "Lê Trọng Chiến", "https://www.facebook.com/100000273215197 (chưa xác nhận trùng tài khoản)",
         "ROXANA - QUYẾT TÂM ĐÒI NHÀ", "2025", "Tóm tắt sự tình dự án Roxana từ đầu đến đuôi.",
         "71 thích, 14 bl, 8 cs", "Thông tin/văn bản tổng hợp",
         "https://www.facebook.com/photo/?fbid=25484736727785418&set=gm.824243813726416&idorvanity=580590514758415",
         None),
        (5, "Nguyen Toan Roxana", "https://www.facebook.com/61590964520879", "ROXANA PLAZA - HÀNH TRÌNH ĐÒI NHÀ",
         "gần đây", "Chung cư Roxanan Plaza... CĐT Tường Phong vẫn không giao nhà.", "5.100 thích, 604 bl, 188 cs",
         "Thông tin chi tiết",
         "https://www.facebook.com/photo/?fbid=122110986465365484&set=gm.1576457697519292&idorvanity=427692059062534",
         None),
    ], 5):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)

    ws = wb.create_sheet("Nguồn báo chí")
    for i, row in enumerate([
        ("Kết luận thanh tra dự án chung cư Roxana Plaza", "CafeF",
         "https://cafef.vn/ket-luan-thanh-tra-du-an-chung-cu-roxana-plaza-188230825094328667.chn"),
        ("Kết luận thanh tra dự án chung cư Roxana Plaza", "Người Lao Động (NLD)",
         "https://nld.com.vn/kinh-te/ket-luan-thanh-tra-du-an-chung-cu-roxana-plaza-20230824181327305.htm"),
    ], 5):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
    wb.save(path)


@pytest.fixture
def legacy_file(tmp_path):
    path = tmp_path / "legacy_src" / "Roxana_cu.xlsx"
    path.parent.mkdir()
    make_legacy(path)
    return path


def test_import_mini(project, legacy_file):
    summary = import_legacy.import_file(project, legacy_file)
    assert summary["containers"]["added"] == 2
    assert summary["sources_facebook"]["added"] == 5
    assert summary["sources_web"]["added"] == 2
    assert summary["events"]["added"] == 4
    assert all(v["invalid"] == 0 for v in summary.values())
    assert (project.root / "legacy" / "Roxana_cu.xlsx").is_file()

    cfg = project.load_config()
    parties = {p["id"]: p for p in cfg["key_parties"]}
    assert parties["tuongphong"]["primary"] and parties["tuongphong"]["label"] == "Tường Phong"
    assert parties["tuongphong"]["keyword_group"] == "tuongphong" and parties["tuongphong"]["legacy"]
    assert parties["ong-hoang-tung"]["primary"] is False
    assert cfg["disclaimer"].startswith("LƯU Ý")
    assert {c["url"] for c in cfg["containers"]} == {"https://www.facebook.com/groups/427692059062534/",
                                                    "https://www.facebook.com/groups/580590514758415/"}

    records, _ = read_records(project.records_path)
    by_ref = {r.get("legacy_ref"): r for r in records}
    containers = {r["url"]: r for r in records if r["record_type"] == "container"}
    main_group = containers["https://www.facebook.com/groups/427692059062534/"]
    assert main_group["member_count"] == 2425 and main_group["topic_dedicated"] and main_group["scan_mode"] == "full"

    p1 = by_ref["Bài viết MXH nổi bật!A5"]
    assert p1["url_kind"] == "container_only" and p1["container_id"] == main_group["id"]
    assert p1["posted_at_precision"] == "relative_estimate" and p1["metrics"]["reactions"] == 10
    assert p1["claim_type"] == "keu_goi" and p1["origin"] == "legacy" and p1["run_id"] == "LEGACY-IMPORT"
    p2 = by_ref["Bài viết MXH nổi bật!A6"]
    assert p2["tone"] == "gay_gat" and p2["url_kind"] == "permalink" and p2["content_type"] == "photo"
    assert p2["container_id"] is None
    p3 = by_ref["Bài viết MXH nổi bật!A7"]
    assert p3["url_kind"] == "none" and p3["url"] == "" and p3["author_url"] is None
    assert p3["author_kind"] == "unknown" and p3["posted_at"] == "2026-07-29T15:38:00+07:00"
    assert p3["posted_at_precision"] == "exact"
    p4 = by_ref["Bài viết MXH nổi bật!A8"]
    assert p4["author_url"] == "https://www.facebook.com/100000273215197" and "chưa xác nhận" in p4["notes"]
    assert p4["posted_at_precision"] == "year"
    assert p4["container_id"] == containers["https://www.facebook.com/groups/580590514758415/"]["id"]
    p5 = by_ref["Bài viết MXH nổi bật!A9"]
    assert p5["importance"] == "cao" and p5["importance_reason"] == "Tương tác ≥ 100"
    assert p5["metrics"] == {"reactions": 5100, "comments": 604, "shares": 188,
                             "counted_at": "2026-08-22T00:00:00+07:00"}
    assert p5["posted_at_precision"] == "unknown"
    assert p5["keywords_matched"] == ["roxana", "tuongphong"] and p5["entities_mentioned"] == ["tuongphong"]

    events = {r["date_raw"]: r for r in records if r["record_type"] == "event"}
    assert events["2017 - 2020"]["date"] == "2017"
    assert events["Công an Bình Dương"]["date"] is None
    assert events["Công an Bình Dương"]["sort_key"] == "2021-01-22"
    assert events["22/8/2026"]["reliability"] == "mxh" and events["22/1/2021"]["reliability"] == "bao_chi"

    web = [r for r in records if r.get("platform") == "web"]
    assert len(web) == 2 and all(r["author_kind"] == "page" and r["url_kind"] == "permalink" for r in web)


def test_import_twice_adds_nothing(project, legacy_file):
    import_legacy.import_file(project, legacy_file)
    before = project.records_path.read_bytes()
    summary = import_legacy.import_file(project, legacy_file)
    assert project.records_path.read_bytes() == before
    assert all(v["added"] == 0 for v in summary.values())
    assert len(project.load_config()["key_parties"]) == 5  # 4 trong config mẫu + ông Hoàng Tùng


@pytest.mark.skipif(not REAL.exists(), reason="file Excel cũ không có trên máy này")
def test_import_real_file_counts(project):
    summary = import_legacy.import_file(project, REAL)
    assert summary["events"]["added"] == 11
    assert summary["sources_facebook"]["added"] == 11
    assert summary["sources_web"]["added"] == 14
    assert summary["containers"]["added"] == 4
    assert all(v["invalid"] == 0 for v in summary.values()), summary
    assert len([p for p in project.load_config()["key_parties"] if p.get("legacy")]) == 12
