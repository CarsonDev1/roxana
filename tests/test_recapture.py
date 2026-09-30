"""Re-capturing an existing post (cleaner screenshots) via rechecks: displays use the newest images."""
from add_record import add_records
from build_excel import S_COMMENTS, S_SOURCES, build
from factories import ev, fb_comment, fb_source
from openpyxl import load_workbook
from view import build_view
from common import read_records

RUN = "RUN-2026-10-01-01"


def _setup(project, make_png):
    cfg = project.load_config()
    [src] = add_records(project, [fb_source([ev(make_png("old_post.png")),
                                              ev(make_png("old_cs.png"), "cscroll", "Bình luận")])], RUN, cfg)
    old_scroll = [p for p in src["evidence_paths"] if "_cscroll_" in p][0]
    [cmt] = add_records(project, [fb_comment(src["id"], old_scroll)], RUN, cfg)
    [chk] = add_records(project, [{"record_type": "recheck", "target_id": src["id"], "status": "active",
                                   "checked_at": "2026-10-02T10:00:00+07:00",
                                   "evidence": [ev(make_png("new_post.png")),
                                                ev(make_png("new_cs.png"), "cscroll", "Bình luận (ảnh gọn)")]}],
                        RUN, cfg)
    new_scroll = [p for p in chk["evidence_paths"] if "_cscroll_" in p][0]
    return cfg, src, cmt, chk, old_scroll, new_scroll


def test_comment_recheck_can_repoint_scroll_refs(project, make_png):
    cfg, src, cmt, chk, old_scroll, new_scroll = _setup(project, make_png)
    [res] = add_records(project, [{"record_type": "recheck", "target_id": cmt["id"], "status": "active",
                                   "checked_at": "2026-10-02T10:01:00+07:00",
                                   "scroll_refs": [{"file": new_scroll, "position": 1}]}], RUN, cfg)
    assert res["status"] == "added", res
    view = build_view(read_records(project.records_path)[0])
    assert view.by_id[cmt["id"]]["scroll_refs"] == [{"file": new_scroll, "position": 1}]


def test_recheck_scroll_refs_must_belong_to_the_comments_post(project, make_png):
    cfg, src, cmt, *_ = _setup(project, make_png)
    [res] = add_records(project, [{"record_type": "recheck", "target_id": cmt["id"], "status": "active",
                                   "checked_at": "2026-10-02T10:01:00+07:00",
                                   "scroll_refs": [{"file": "screenshots/facebook/x.png", "position": 1}]}], RUN, cfg)
    assert res["status"] == "invalid"


def test_displays_prefer_newest_screenshot(project, make_png):
    cfg, src, cmt, chk, old_scroll, new_scroll = _setup(project, make_png)
    add_records(project, [{"record_type": "recheck", "target_id": cmt["id"], "status": "active",
                           "checked_at": "2026-10-02T10:01:00+07:00",
                           "scroll_refs": [{"file": new_scroll, "position": 1}]}], RUN, cfg)
    wb = load_workbook(build(project)["file"])
    ws = wb[S_SOURCES]
    h = {c.value: i + 1 for i, c in enumerate(ws[1])}
    assert ws.cell(2, h["Mở ảnh gốc"]).hyperlink.target == "../" + chk["evidence_paths"][0]
    wc = wb[S_COMMENTS]
    hc = {c.value: i + 1 for i, c in enumerate(wc[1])}
    assert wc.cell(2, hc["Ảnh cuộn"]).hyperlink.target == "../" + new_scroll
