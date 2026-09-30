import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

import build_excel
from add_record import add_records
from build_excel import (S_CHANGES, S_COMMENTS, S_CONTAINERS, S_EVIDENCE, S_EXCLUDED, S_LEGEND, S_LOG,
                         S_OVERVIEW, S_PEOPLE, S_SOURCES, S_TIMELINE, SHEET_ORDER, ExcelLockedError, build)
from factories import ev, fb_comment, fb_container, fb_source
from formula_eval import eval_formula

RUN = "RUN-2026-09-29-01"


@pytest.fixture
def populated(project, make_png):
    cfg = project.load_config()

    def add(recs):
        return add_records(project, recs, RUN, cfg)

    [ctr] = add([fb_container()])
    [src] = add([fb_source([ev(make_png("p.png")), ev(make_png("cs1.png"), "cscroll", "Bình luận 1-3")],
                           container_id=ctr["id"])])
    scroll = [p for p in src["evidence_paths"] if "_cscroll_" in p][0]
    comments = add([
        fb_comment(src["id"], scroll, 1, text="=))) trả nhà đi"),
        fb_comment(src["id"], scroll, 2, text="Bà Phạm Thị Ngọc Liên phải chịu trách nhiệm", depth=2,
                   parent_comment_id="@0", entities_mentioned=["lien"], tone="gay_gat", claim_type="cao_buoc",
                   importance="cao", importance_reason="Nêu đích danh bên chính kèm cáo buộc",
                   evidence=[ev(make_png("c2.png"), "comment", "Bình luận")]),
        fb_comment(src["id"], scroll, 3, text="Ký tự lạ\x0b\x00 ở đây"),
    ])
    legacy = fb_source([], origin="legacy", url="https://www.facebook.com/groups/472914441060049/",
                       url_kind="container_only", author_name="Huynh Bich Diem",
                       author_url="https://www.facebook.com/100009329476350",
                       text="Nơi hội tụ tinh hoa???? Xạo, trả nhà cho người dân", tone="gay_gat", claim_type=None,
                       importance="trung_binh", importance_reason=None, entities_mentioned=[],
                       posted_at="2024-12-02T00:00:00+07:00", posted_at_precision="day", posted_at_raw="2/12/2024")
    legacy.pop("snapshot_text")
    add([legacy])
    add([
        {"record_type": "search_log", "platform": "facebook", "scope": "global", "section": "posts",
         "query": "Roxana Plaza", "filters": {"year": 2021}, "started_at": "2026-09-29T20:00:00+07:00",
         "ended_at": "2026-09-29T20:10:00+07:00", "results_seen": 40, "results_new": 1, "results_duplicate": 0,
         "results_excluded": 1, "reached_end": False, "issues": "Facebook chỉ trả 40 kết quả"},
        {"record_type": "exclusion", "url": "https://www.facebook.com/somebody/posts/1",
         "excerpt": "Chúc mừng sinh nhật chị Phạm Ngọc Liên", "keywords_matched": ["lien"],
         "reason": "Trùng tên — không có từ ngữ cảnh"},
        {"record_type": "recheck", "target_id": src["id"], "checked_at": "2026-10-05T20:00:00+07:00",
         "status": "edited", "new_text": "Nội dung đã sửa", "metrics": {"reactions": 50}},
        {"record_type": "event", "date_raw": "22/8/2026", "date": "2026-08-22",
         "description": "Buổi đối thoại tại Tiếp công dân TP.HCM", "source_text": "Giấy mời",
         "related_ids": [src["id"]], "reliability": "van_ban"},
    ])
    return project, src, comments


def load(project):
    result = build(project)
    return result, load_workbook(result["file"])


def header_map(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}


def row_of(ws, id_):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == id_:
            return r
    raise AssertionError(f"{id_} not found in {ws.title}")


def test_all_sheets_in_order(populated):
    project, *_ = populated
    _, wb = load(project)
    assert wb.sheetnames == SHEET_ORDER


def test_sources_sheet(populated):
    project, src, _ = populated
    _, wb = load(project)
    ws = wb[S_SOURCES]
    h = header_map(ws)
    assert ws.max_row == 3  # tiêu đề + bài quét + bài từ file cũ
    r = row_of(ws, src["id"])
    assert ws.cell(r, h["Trạng thái"]).value == "Đã sửa"
    assert ws.cell(r, h["Thích"]).value == 50
    assert ws.cell(r, h["Nội dung nguyên văn"]).value.startswith("Hôm qua")  # giữ nội dung gốc
    assert ws.cell(r, h["Mở ảnh gốc"]).hyperlink.target == "../" + src["evidence_paths"][0]
    assert ws.cell(r, h["Bình luận (đã thu)"]).value.startswith("=COUNTIF(")
    assert eval_formula(wb, ws.cell(r, h["Bình luận (đã thu)"]).value) == 3
    legacy_r = row_of(ws, "FB-P00002")
    assert ws.cell(legacy_r, h["Ảnh"]).value == "Chưa có ảnh"
    assert ws.cell(legacy_r, h["Nguồn dữ liệu"]).value == "Từ file cũ"
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref.startswith("A1:")
    assert ws.cell(1, 1).font.name == "Arial"


def test_thumbnails_anchored(populated):
    project, src, _ = populated
    _, wb = load(project)
    ws = wb[S_SOURCES]
    anchors = {(img.anchor._from.row + 1, img.anchor._from.col + 1) for img in ws._images}
    assert (row_of(ws, src["id"]), header_map(ws)["Ảnh"]) in anchors
    assert list((project.root / "output" / "thumbs").glob("*.jpg"))


def test_comments_sheet_links_and_text_safety(populated):
    project, src, comments = populated
    _, wb = load(project)
    ws = wb[S_COMMENTS]
    h = header_map(ws)
    assert ws.max_row == 4
    r0 = row_of(ws, comments[0]["id"])
    cell = ws.cell(r0, h["Nội dung nguyên văn"])
    assert cell.value == "=))) trả nhà đi" and cell.data_type == "s"
    assert ws.cell(r0, h["Thuộc bài"]).hyperlink.location == f"'{S_SOURCES}'!A{row_of(wb[S_SOURCES], src['id'])}"
    r1 = row_of(ws, comments[1]["id"])
    assert ws.cell(r1, h["Trả lời cho"]).hyperlink.location == f"'{S_COMMENTS}'!A{r0}"
    assert "vị trí 2" in ws.cell(r1, h["Ảnh cuộn"]).value
    r2 = row_of(ws, comments[2]["id"])
    assert ws.cell(r2, h["Nội dung nguyên văn"]).value == "Ký tự lạ ở đây"


def test_very_long_text_is_truncated_with_note(project, make_png):
    add_records(project, [fb_source([ev(make_png())], text="a" * 40000)], RUN, project.load_config())
    _, wb = load(project)
    ws = wb[S_SOURCES]
    value = ws.cell(2, header_map(ws)["Nội dung nguyên văn"]).value
    assert len(value) <= 32767 and "bị cắt" in value


def test_containers_sheet_flags_join_needed(populated):
    project, *_ = populated
    _, wb = load(project)
    ws = wb[S_CONTAINERS]
    h = header_map(ws)
    assert ws.cell(2, h["Cần xin vào"]).value == "Có"
    assert ws.cell(2, 1).fill.start_color.rgb.endswith("FFFF00")
    assert eval_formula(wb, ws.cell(2, h["Số bài thu được"]).value) == 1


def test_log_changes_excluded_sheets(populated):
    project, *_ = populated
    _, wb = load(project)
    log = wb[S_LOG]
    hl = header_map(log)
    assert log.cell(2, hl["Hết kết quả"]).value == "Không"
    assert "40" in log.cell(2, hl["Sự cố/giới hạn"]).value
    ch = wb[S_CHANGES]
    hc = header_map(ch)
    assert ch.max_row == 2
    assert ch.cell(2, hc["Trạng thái"]).value == "Đã sửa"
    assert ch.cell(2, hc["Nội dung mới"]).value == "Nội dung đã sửa"
    ex = wb[S_EXCLUDED]
    he = header_map(ex)
    assert ex.max_row == 2 and "Trùng tên" in ex.cell(2, he["Lý do"]).value
    assert "Người đăng" not in he


def test_empty_store_builds(project):
    result, wb = load(project)
    assert result["status"] == "ok" and wb[S_SOURCES].max_row == 1


@pytest.mark.skipif(sys.platform != "win32", reason="khoá file kiểu Windows")
def test_locked_output_reports_error(populated):
    project, *_ = populated
    out = Path(build(project)["file"])
    with open(out, "rb"):
        with pytest.raises(ExcelLockedError):
            build(project)
    assert not list(out.parent.glob("*.tmp.xlsx"))


def test_cli_prints_json(populated, capsys):
    project, *_ = populated
    assert build_excel.main(["--project", str(project.root)]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "ok"


def test_overview_formulas_evaluate_to_expected(populated):
    project, *_ = populated
    _, wb = load(project)
    ws = wb[S_OVERVIEW]
    values = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}

    def value(label_text):
        return eval_formula(wb, values[label_text])

    assert value("Bài viết & nguồn") == 2
    assert value("Bình luận") == 3
    assert value("Nhóm & trang") == 1
    assert value("Nhóm kín cần xin vào") == 1
    assert value("Kết quả đã loại trừ (trùng tên)") == 1
    assert value("Facebook") == 2
    assert value("Web / Báo chí") == 0
    assert value("Bà Phạm Thị Ngọc Liên") == 2   # bài FB-P00001 + bình luận thứ 2
    assert value("Gay gắt") == 2                  # bài từ file cũ + bình luận thứ 2
    assert value("Cao") == 2
    assert value("2026-08") == 1 and value("2024-12") == 1 and value("2026-09") == 3
    assert value("Không rõ") == 0
    assert any("Facebook chỉ trả 40 kết quả" == ws.cell(r, 3).value for r in range(1, ws.max_row + 1))


def test_people_sheet(populated):
    project, *_ = populated
    _, wb = load(project)
    ws = wb[S_PEOPLE]
    h = header_map(ws)
    names = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    assert names[:4] == ["Công ty TNHH XD-DV-TM-Đầu tư BĐS Tường Phong", "Công ty CP Naviland",
                         "Bà Phạm Thị Ngọc Liên", "Toà án Nhân dân Khu vực 16 – TP.HCM"]
    lien = names.index("Bà Phạm Thị Ngọc Liên") + 2
    assert ws.cell(lien, h["Nhóm"]).value == "Bên chính"
    assert eval_formula(wb, ws.cell(lien, h["Số lần được nhắc"]).value) == 2
    qc = names.index("Quyết Chiến Roxana") + 2
    assert eval_formula(wb, ws.cell(qc, h["Số bài đăng"]).value) == 1
    hbd = names.index("Huynh Bich Diem") + 2
    assert names.count("Huynh Bich Diem") == 1      # gộp theo link trang cá nhân
    assert eval_formula(wb, ws.cell(hbd, h["Số bài đăng"]).value) == 1
    assert eval_formula(wb, ws.cell(hbd, h["Số bình luận"]).value) == 3


def test_evidence_sheet_lists_high_importance(populated):
    project, src, comments = populated
    _, wb = load(project)
    ws = wb[S_EVIDENCE]
    h = header_map(ws)
    assert [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)] == [src["id"], comments[1]["id"]]
    assert len(ws._images) == 2
    assert len(ws.cell(2, h["SHA-256"]).value) == 64


def test_timeline_sheet(populated):
    project, *_ = populated
    _, wb = load(project)
    ws = wb[S_TIMELINE]
    h = header_map(ws)
    assert ws.cell(2, h["Độ tin cậy"]).value == "Có văn bản chính thức"
    assert ws.cell(2, h["Mã liên quan"]).hyperlink.location.startswith(f"'{S_SOURCES}'!A")


def test_legend_has_definitions_and_disclaimer(populated):
    project, *_ = populated
    _, wb = load(project)
    ws = wb[S_LEGEND]
    text = "\n".join(str(ws.cell(r, c).value or "") for r in range(1, ws.max_row + 1) for c in (1, 2))
    for needle in ("Gay gắt", "Cáo buộc", "certutil -hashfile", "không phải kết luận pháp lý", "chép cả thư mục"):
        assert needle in text


def test_full_calc_on_load_and_idempotent(populated):
    project, *_ = populated

    def snapshot():
        wb = load_workbook(build(project)["file"])
        cells = {n: [[c.value for c in row] for row in wb[n].iter_rows()] for n in SHEET_ORDER if n != S_OVERVIEW}
        return wb, cells

    wb1, first = snapshot()
    _, second = snapshot()
    assert wb1.calculation.fullCalcOnLoad is True
    assert first == second
