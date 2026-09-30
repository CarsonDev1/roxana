"""Readability / accessibility styling of the workbook (user feedback after the first legacy build)."""
from openpyxl import load_workbook

from build_excel import (S_COMMENTS, S_CONTAINERS, S_LEGEND, S_OVERVIEW, S_SOURCES, SHEET_ORDER, build)
from test_build_excel import header_map, populated, row_of  # noqa: F401  (fixture re-export)


def _wb(project):
    return load_workbook(build(project)["file"])


def _rgb(color):
    return (color.rgb or "")[-6:] if color is not None else ""


def test_header_row_is_high_contrast(populated):
    project, *_ = populated
    ws = _wb(project)[S_SOURCES]
    cell = ws.cell(1, 1)
    assert _rgb(cell.fill.start_color) == "1F3864"
    assert cell.font.bold and _rgb(cell.font.color) == "FFFFFF"
    assert ws.row_dimensions[1].height and ws.row_dimensions[1].height >= 30


def test_body_font_is_readable_size(populated):
    project, src, _ = populated
    ws = _wb(project)[S_SOURCES]
    r = row_of(ws, src["id"])
    assert ws.cell(r, header_map(ws)["Nội dung nguyên văn"]).font.size >= 11


def test_rows_are_banded_and_bordered(populated):
    project, *_ = populated
    ws = _wb(project)[S_COMMENTS]
    col = header_map(ws)["Người viết"]
    fills = {_rgb(ws.cell(r, col).fill.start_color) for r in (2, 3)}
    assert "F2F2F2" in fills and len(fills) == 2  # dòng xen kẽ
    assert ws.cell(2, col).border.bottom.style == "thin"


def test_high_importance_and_harsh_tone_are_highlighted(populated):
    project, _, comments = populated
    ws = _wb(project)[S_COMMENTS]
    h = header_map(ws)
    r = row_of(ws, comments[1]["id"])
    for header in ("Mức quan trọng", "Thái độ"):
        cell = ws.cell(r, h[header])
        assert _rgb(cell.fill.start_color) == "FCE4D6", header
        assert cell.value in ("Cao", "Gay gắt")  # màu luôn đi kèm chữ


def test_changed_status_is_highlighted(populated):
    project, src, _ = populated
    ws = _wb(project)[S_SOURCES]
    cell = ws.cell(row_of(ws, src["id"]), header_map(ws)["Trạng thái"])
    assert cell.value == "Đã sửa" and _rgb(cell.fill.start_color) == "FFF2CC"


def test_join_needed_row_keeps_yellow_over_banding(populated):
    project, *_ = populated
    ws = _wb(project)[S_CONTAINERS]
    assert _rgb(ws.cell(2, 2).fill.start_color) == "FFFF00"


def test_overview_has_table_of_contents_with_links(populated):
    project, *_ = populated
    wb = _wb(project)
    ws = wb[S_OVERVIEW]
    targets = {c.hyperlink.location for row in ws.iter_rows() for c in row if c.hyperlink and c.hyperlink.location}
    for name in SHEET_ORDER[1:]:
        assert f"'{name}'!A1" in targets, name


def test_sheets_have_tab_colors_and_print_setup(populated):
    project, *_ = populated
    wb = _wb(project)
    for name in SHEET_ORDER:
        ws = wb[name]
        assert ws.sheet_properties.tabColor is not None, name
    ws = wb[S_SOURCES]
    assert ws.page_setup.orientation == "landscape"
    assert ws.sheet_properties.pageSetUpPr.fitToPage and ws.page_setup.fitToWidth == 1
    assert ws.print_title_rows.replace("$", "") == "1:1"


def test_legend_section_titles_are_styled(populated):
    project, *_ = populated
    ws = _wb(project)[S_LEGEND]
    titles = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Thái độ"]
    assert titles and ws.cell(titles[0], 1).font.bold
