"""Inputs from the whole-branch review that must never break the append-only store or the workbook."""
import zipfile
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

import build_excel
import xlsx_helpers
from add_record import add_records
from build_excel import S_COMMENTS, build
from common import append_records, read_records
from factories import ev, fb_comment, fb_source

RUN = "RUN-2026-10-01-01"


def _add(project, recs):
    return add_records(project, recs, RUN, project.load_config())


def _source(project, make_png, n=1, **kw):
    [res] = _add(project, [fb_source([ev(make_png(f"p{n}.png")), ev(make_png(f"cs{n}.png"), "cscroll", "Bình luận")],
                                     url=f"https://www.facebook.com/groups/1/posts/{n}/", **kw)])
    assert res["status"] == "added", res
    return res


def _sheets_parse(path):
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                ET.fromstring(z.read(name))  # raises on XML Excel cannot open


@pytest.mark.parametrize("content", [b"", b"%PDF-1.4 not an image"])
def test_non_image_evidence_is_invalid(project, tmp_path, content):
    bad = tmp_path / "incoming" / "bad.png"
    bad.parent.mkdir(parents=True, exist_ok=True)
    bad.write_bytes(content)
    [res] = _add(project, [fb_source([ev(bad)], url="https://www.facebook.com/groups/1/posts/77/")])
    assert res["status"] == "invalid" and any("không phải ảnh" in e for e in res["errors"]), res
    assert not project.records_path.exists()


def test_build_survives_unreadable_image_already_in_store(project, make_png):
    res = _source(project, make_png)
    records, _ = read_records(project.records_path)
    broken = project.resolve(res["evidence_paths"][0])
    broken.write_bytes(b"")  # ảnh hỏng sau khi ghi (ổ đĩa, người dùng…)
    result = build(project)
    assert result["status"] == "ok"
    assert any(broken.name in w for w in result["warnings"]), result["warnings"]


def test_noncharacters_are_stripped_so_excel_can_open(project, make_png):
    src = _source(project, make_png)
    scroll = [p for p in src["evidence_paths"] if "_cscroll_" in p][0]
    [res] = _add(project, [fb_comment(src["id"], scroll, text="abc￾￿d\ud83d",
                                      author_url="https://www.facebook.com/1￿")])
    assert res["status"] == "added", res
    path = build(project)["file"]
    _sheets_parse(path)
    ws = load_workbook(path)[S_COMMENTS]
    assert any(ws.cell(r, c).value == "abcd" for r in range(2, ws.max_row + 1) for c in range(1, ws.max_column + 1))


def test_hyperlinks_per_sheet_are_capped_with_warning(project, make_png, monkeypatch):
    monkeypatch.setattr(xlsx_helpers, "MAX_HYPERLINKS", 6)
    src = _source(project, make_png)
    scroll = [p for p in src["evidence_paths"] if "_cscroll_" in p][0]
    _add(project, [fb_comment(src["id"], scroll, i + 1, text=f"bình luận {i}", url=f"https://www.facebook.com/c/{i}")
                   for i in range(5)])
    result = build(project)
    ws = load_workbook(result["file"])[S_COMMENTS]
    links = sum(1 for row in ws.iter_rows() for c in row if c.hyperlink)
    assert links <= 6
    assert any("65.530" in w or "hyperlink" in w for w in result["warnings"]), result["warnings"]


def test_author_url_tracking_params_are_stripped(project, make_png):
    raw = ("https://www.facebook.com/groups/427692059062534/user/100093187612816/?__cft__[0]=" + "AZ" * 150
           + "&__tn__=%2CO%2CP-R")
    res = _source(project, make_png, author_url=raw)
    records, _ = read_records(project.records_path)
    stored = next(r for r in records if r["id"] == res["id"])
    assert stored["author_url"] == "https://www.facebook.com/groups/427692059062534/user/100093187612816"


def test_profile_php_keeps_id_and_drops_comment_id(project, make_png):
    res = _source(project, make_png, author_url="https://m.facebook.com/profile.php?id=42&comment_id=9&__tn__=R")
    records, _ = read_records(project.records_path)
    assert next(r for r in records if r["id"] == res["id"])["author_url"] == \
        "https://www.facebook.com/profile.php?id=42"


@pytest.mark.parametrize("patch", [
    {"evidence": ["a.png"]},
    {"evidence": {"file": "a.png", "kind": "post", "shows": "x"}},
    {"scroll_refs": ["a.png"]},
    {"attachments": ["ảnh"]},
])
def test_malformed_shapes_are_invalid_not_crash(project, make_png, patch):
    src = _source(project, make_png)
    scroll = [p for p in src["evidence_paths"] if "_cscroll_" in p][0]
    rec = fb_comment(src["id"], scroll)
    rec.update(patch)
    [res] = _add(project, [rec])
    assert res["status"] == "invalid", res


def test_non_object_batch_element_is_invalid(project):
    results = _add(project, ["hello", 42])
    assert [r["status"] for r in results] == ["invalid", "invalid"]


@pytest.mark.parametrize("extra", [
    {"evidence": [{"file": "x.png", "kind": "post", "shows": "bài"}]},
    {"snapshot_text": "Phạm Ngọc Liên · SĐT 09xxxx"},
    {"author": "Phạm Ngọc Liên"},
    {"profile_url": "https://www.facebook.com/someone"},
    {"shared_from": {"author_name": "Phạm Ngọc Liên"}},
])
def test_exclusion_only_accepts_allowlisted_fields(project, extra):
    rec = {"record_type": "exclusion", "url": "https://www.facebook.com/somebody/posts/1",
           "excerpt": "Chúc mừng sinh nhật chị Phạm Ngọc Liên", "keywords_matched": ["lien"],
           "reason": "Trùng tên — không có từ ngữ cảnh", **extra}
    [res] = _add(project, [rec])
    assert res["status"] == "invalid", res
    assert not project.records_path.exists()
    assert not (project.root / "screenshots").exists() and not (project.root / "snapshots").exists()
