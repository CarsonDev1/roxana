"""openpyxl helpers: fonts, table writer, thumbnails, hyperlinks and cell-value sanitising."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from PIL import Image

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE_FONT = Font(name="Arial", size=14, bold=True)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
HEADER_FILL = PatternFill("solid", start_color="D9D9D9")
YELLOW_FILL = PatternFill("solid", start_color="FFFF00")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
MAX_CELL_CHARS = 32767
MAX_THUMB_HEIGHT = 540  # px — giữ chiều cao hàng dưới giới hạn 409pt của Excel
TRUNCATION_NOTE = " …[bị cắt do giới hạn 32.767 ký tự của ô Excel — xem bản chữ gốc]"
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass
class Img:
    path: Path
    width: int = 240


@dataclass
class Link:
    target: str
    text: str


@dataclass
class Internal:
    sheet: str
    row: int
    text: str


@dataclass
class Formula:
    expr: str


@dataclass
class Column:
    header: str
    width: float
    get: Callable[[dict], Any]


def safe_text(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = _ILLEGAL.sub("", str(value))
    if len(text) > MAX_CELL_CHARS:
        text = text[: MAX_CELL_CHARS - len(TRUNCATION_NOTE)] + TRUNCATION_NOTE
    return text


def write_cell(ws, row: int, col: int, value) -> None:
    cell = ws.cell(row=row, column=col)
    cell.alignment = WRAP_TOP
    if isinstance(value, Formula):
        cell.value = value.expr
        cell.font = FONT
        return
    if isinstance(value, Link):
        cell.hyperlink = value.target
        text, font = value.text, LINK_FONT
    elif isinstance(value, Internal):
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{value.sheet}'!A{value.row}")
        text, font = value.text, LINK_FONT
    else:
        text, font = value, FONT
    cell.value = safe_text(text)
    if isinstance(cell.value, str):
        cell.data_type = "s"  # bình luận như "=)))" phải là chữ, không bao giờ thành công thức
    cell.font = font


def make_thumbnail(src: Path, thumbs_dir: Path, width: int) -> Path:
    thumbs_dir.mkdir(parents=True, exist_ok=True)
    dest = thumbs_dir / f"{src.stem}_w{width}.jpg"
    if dest.exists():
        return dest
    with Image.open(src) as im:
        im = im.convert("RGB")
        height = max(1, round(im.height * width / im.width))
        im = im.resize((width, height))
        max_height = min(int(width * 1.5), MAX_THUMB_HEIGHT)
        if height > max_height:
            im = im.crop((0, 0, width, max_height))
        im.save(dest, "JPEG", quality=80)
    return dest


def add_thumbnail(ws, anchor: str, src: Path, thumbs_dir: Path, width: int) -> int:
    """Embed a thumbnail at `anchor`; return its height in px (0 when the original is missing)."""
    if not src.is_file():
        return 0
    img = XLImage(str(make_thumbnail(src, thumbs_dir, width)))
    ws.add_image(img, anchor)
    return img.height


def write_table(ws, columns: list[Column], items: list[dict], thumbs_dir: Path,
                fill: Callable[[dict], PatternFill | None] | None = None) -> None:
    for c, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=col.header)
        cell.font, cell.fill, cell.alignment = BOLD, HEADER_FILL, WRAP_TOP
        ws.column_dimensions[get_column_letter(c)].width = col.width
    for r, item in enumerate(items, 2):
        tallest = 0
        for c, col in enumerate(columns, 1):
            value = col.get(item)
            if isinstance(value, Img):
                tallest = max(tallest, add_thumbnail(ws, f"{get_column_letter(c)}{r}", value.path, thumbs_dir,
                                                     value.width))
                ws.cell(row=r, column=c).font = FONT
            else:
                write_cell(ws, r, c, value)
        row_fill = fill(item) if fill else None
        if row_fill:
            for c in range(1, len(columns) + 1):
                ws.cell(row=r, column=c).fill = row_fill
        if tallest:
            ws.row_dimensions[r].height = tallest * 0.75 + 4
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(items) + 1}"


def col_letter(columns: list[Column], header: str) -> str:
    return get_column_letter([c.header for c in columns].index(header) + 1)


def countif_literal(value: str) -> str:
    """Escape a value so COUNTIF matches it literally inside a formula string."""
    return value.replace("~", "~~").replace("*", "~*").replace("?", "~?").replace('"', '""')


def countif(sheet: str, letter: str, criterion: str) -> str:
    return f"COUNTIF('{sheet}'!{letter}:{letter},\"{criterion}\")"
