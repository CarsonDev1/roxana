"""One-off import of the legacy workbook (Roxana_Plaza_Tong_hop_vu_viec.xlsx) into config.json and the store.

Usage: python import_legacy.py [--project D:\\Roxana] --file "C:\\...\\Roxana_Plaza_Tong_hop_vu_viec (1) (1).xlsx"
Safe to run twice: the second run adds nothing.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

from add_record import add_records
from common import Project, fold, match_keywords, normalize_url, setup_stdout

LEGACY_RUN = "LEGACY-IMPORT"
LEGACY_AT = "2026-08-22T00:00:00+07:00"  # ngày tổng hợp ghi trong file cũ
SHEET_EVENTS = "Tóm tắt vụ việc"
SHEET_PARTIES = "Các bên liên quan"
SHEET_POSTS = "Bài viết MXH nổi bật"
SHEET_PRESS = "Nguồn báo chí"
FIRST_ROW = 5
PRIMARY_PARTIES = {
    "Công ty TNHH XD-DV-TM-Đầu tư BĐS Tường Phong": ("tuongphong", "Tường Phong"),
    "Công ty CP Naviland": ("naviland", "Naviland"),
    "Công ty CP Đầu tư Viethome": ("viethome", "Viethome"),
    "Ông Lầu Nam Tường": ("launamtuong", "Ông Lầu Nam Tường"),
    "Bà Phạm Thị Ngọc Liên": ("lien", "Bà Phạm Thị Ngọc Liên"),
    "Bà Dương Thị Phương Tuyền": ("tuyen", "Bà Dương Thị Phương Tuyền"),
}
_DMY = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
_HM = re.compile(r"(\d{1,2}):(\d{2})")
_YEAR = re.compile(r"\d{4}")
_GROUP = re.compile(r"/groups/(\d+)")
_VANITY = re.compile(r"idorvanity=(\d+)")
_MEMBERS = re.compile(r"\s*\(([\d.]+)\s*thành viên\)")


def _cell(ws, row: int, col: int) -> str:
    value = ws.cell(row, col).value
    return str(value).strip() if value is not None else ""


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", fold(text)).strip("-")[:48]


def parse_time(raw: str) -> tuple[str | None, str]:
    text = raw.strip()
    m = _DMY.search(text)
    if m:
        day, month, year = (int(x) for x in m.groups())
        date = f"{year:04d}-{month:02d}-{day:02d}"
        hm = _HM.search(text[m.end():])
        if hm:
            return f"{date}T{int(hm[1]):02d}:{hm[2]}:00+07:00", "exact"
        return f"{date}T00:00:00+07:00", ("relative_estimate" if text.startswith("~") else "day")
    if re.fullmatch(r"\d{4}", text):
        return f"{text}-01-01T00:00:00+07:00", "year"
    return None, "unknown"


def event_date(raw: str) -> str | None:
    m = _DMY.search(raw)
    if m:
        day, month, year = (int(x) for x in m.groups())
        return f"{year:04d}-{month:02d}-{day:02d}"
    y = _YEAR.search(raw)
    return y[0] if y else None


def reliability(source_text: str) -> str:
    f = fold(source_text)
    return "mxh" if ("facebook" in f or f.startswith("bai dang") or f.startswith("anh ")) else "bao_chi"


def parse_metrics(raw: str) -> dict:
    metrics = {}
    for key, word in (("reactions", "thích"), ("comments", "bl"), ("shares", "cs")):
        m = re.search(r"([\d.]+)\s*" + word, raw)
        if m:
            metrics[key] = int(m[1].replace(".", ""))
    if metrics:
        metrics["counted_at"] = LEGACY_AT
    return metrics


def author_url(raw: str) -> str | None:
    token = raw.split()[0] if raw else ""
    return token if token.startswith("http") else None


def container_name(raw: str) -> str:
    name = _MEMBERS.sub("", raw).strip()
    if name.startswith("chia sẻ lại") and "nhóm " in name:
        name = name.split("nhóm ", 1)[1].strip()
    return name


def group_url(link: str) -> str | None:
    m = _GROUP.search(link) or _VANITY.search(link)
    return f"https://www.facebook.com/groups/{m[1]}/" if m else None


def _post_rows(ws):
    for row in range(FIRST_ROW, ws.max_row + 1):
        if _cell(ws, row, 2):
            yield row


def _name_urls(ws) -> dict[str, str]:
    urls: dict[str, str] = {}
    for row in _post_rows(ws):
        url = group_url(_cell(ws, row, 9))
        if url:
            urls.setdefault(fold(container_name(_cell(ws, row, 4))), url)
    return urls


def merge_parties(config: dict, ws) -> None:
    groups = {g["id"] for g in config.get("keyword_groups", [])}
    parties = config.setdefault("key_parties", [])
    by_id = {p["id"]: p for p in parties}
    for row in range(FIRST_ROW, ws.max_row + 1):
        name = _cell(ws, row, 1)
        if not name:
            continue
        primary = name in PRIMARY_PARTIES
        pid, short = PRIMARY_PARTIES.get(name, (slug(name), name))
        entry = {"id": pid, "name": name, "label": short, "kind": _cell(ws, row, 2), "role": _cell(ws, row, 3),
                 "press_sources": _cell(ws, row, 4), "keyword_group": pid if primary and pid in groups else None,
                 "primary": primary, "legacy": True}
        if pid in by_id:
            by_id[pid].update(entry)
        else:
            parties.append(entry)
            by_id[pid] = entry


def event_records(ws) -> tuple[list[dict], str | None]:
    records, disclaimer, sort_key = [], None, ""
    for row in range(FIRST_ROW, ws.max_row + 1):
        a, b, c = _cell(ws, row, 1), _cell(ws, row, 2), _cell(ws, row, 3)
        if a.upper().startswith("LƯU Ý"):
            disclaimer = a
            continue
        if not (a and b):
            continue
        date = event_date(a)
        sort_key = max(date or "", sort_key)
        records.append({"record_type": "event", "origin": "legacy", "date_raw": a, "date": date,
                        "sort_key": sort_key, "description": b, "source_text": c, "related_ids": [],
                        "reliability": reliability(c), "legacy_ref": f"{SHEET_EVENTS}!A{row}",
                        "captured_at": LEGACY_AT})
    return records, disclaimer


def container_records(ws, config: dict) -> list[dict]:
    name_urls = _name_urls(ws)
    found: dict[str, dict] = {}
    for row in _post_rows(ws):
        raw = _cell(ws, row, 4)
        name = container_name(raw)
        url = group_url(_cell(ws, row, 9)) or name_urls.get(fold(name))
        if not url:
            continue
        info = found.setdefault(url, {"name": name, "members": None})
        members = _MEMBERS.search(raw)
        if members:
            info["members"] = int(members[1].replace(".", ""))
    records = []
    for url, info in found.items():
        dedicated = any(fold(t) in fold(info["name"]) for t in config.get("context_terms", []))
        records.append({"record_type": "container", "origin": "legacy", "platform": "facebook", "kind": "group",
                        "name": info["name"], "url": url, "privacy": "unknown", "joined": "unknown",
                        "topic_dedicated": dedicated, "scan_mode": "full" if dedicated else "keyword",
                        "member_count": info["members"], "member_count_at": LEGACY_AT if info["members"] else None,
                        "notes": "Từ file cũ", "captured_at": LEGACY_AT})
    return records


def merge_containers(config: dict, containers: list[dict]) -> None:
    listed = config.setdefault("containers", [])
    known = {normalize_url(c["url"]) for c in listed}
    for c in containers:
        if normalize_url(c["url"]) not in known:
            listed.append({"platform": c["platform"], "name": c["name"], "url": c["url"],
                           "privacy": c["privacy"], "joined": c["joined"], "scan_mode": c["scan_mode"]})
            known.add(normalize_url(c["url"]))


def _legacy_tone(level: str) -> str | None:
    return "gay_gat" if "GAY GẮT" in level.upper() else None


def _legacy_claim(level: str) -> str | None:
    f = fold(level)
    if "keu goi" in f:
        return "keu_goi"
    if "phap ly" in f or "van ban" in f:
        return "van_ban"
    if "thong tin" in f:
        return "thong_tin"
    return None


def post_records(ws, config: dict, container_ids: dict[str, str]) -> list[dict]:
    name_urls = _name_urls(ws)
    party_groups = {p["id"]: p.get("keyword_group") for p in config.get("key_parties", [])}
    records = []
    for row in _post_rows(ws):
        author, author_raw, group_raw, when, text, interactions, level, link, note = (
            _cell(ws, row, col) for col in range(2, 11))
        name = container_name(group_raw)
        gurl = group_url(link) or name_urls.get(fold(name))
        if link.startswith("http"):
            url = link
            url_kind = "container_only" if gurl and normalize_url(link) == normalize_url(gurl) else "permalink"
        else:
            url, url_kind = (gurl, "container_only") if gurl else ("", "none")
        posted_at, precision = parse_time(when)
        metrics = parse_metrics(interactions)
        a_url = author_url(author_raw)
        matched, _ = match_keywords(f"{text} {name}", config)
        total = sum(metrics.get(k, 0) for k in ("reactions", "comments", "shares"))
        if "QUAN TRỌNG" in note.upper():
            importance, reason = "cao", "Ghi chú file cũ: bài quan trọng"
        elif total >= 100:
            importance, reason = "cao", "Tương tác ≥ 100"
        else:
            importance, reason = "trung_binh", None
        notes = [f"Mức độ (file cũ): {level}"] if level else []
        if note:
            notes.append(note)
        if author_raw and author_raw != a_url:
            notes.append(f"Link người đăng (file cũ): {author_raw}")
        content_type = ("photo" if "/photo" in link
                        else "shared_post" if group_raw.startswith("chia sẻ lại") else "post")
        records.append({
            "record_type": "source", "origin": "legacy", "platform": "facebook", "content_type": content_type,
            "url": url, "url_kind": url_kind,
            "container_id": container_ids.get(normalize_url(gurl)) if gurl else None,
            "container_name": name, "author_name": author, "author_url": a_url,
            "author_kind": "person" if a_url else "unknown", "posted_at_raw": when, "posted_at": posted_at,
            "posted_at_precision": precision, "text": text, "attachments": [], "metrics": metrics,
            "keywords_matched": matched,
            "entities_mentioned": [pid for pid, g in party_groups.items() if g and g in matched],
            "topics": [], "tone": _legacy_tone(level), "claim_type": _legacy_claim(level),
            "importance": importance, "importance_reason": reason, "notes": " | ".join(notes),
            "legacy_ref": f"{SHEET_POSTS}!A{row}", "captured_at": LEGACY_AT,
        })
    return records


def press_records(ws, config: dict) -> list[dict]:
    party_groups = {p["id"]: p.get("keyword_group") for p in config.get("key_parties", [])}
    records = []
    for row in range(FIRST_ROW, ws.max_row + 1):
        title, outlet, link = _cell(ws, row, 1), _cell(ws, row, 2), _cell(ws, row, 3)
        if not title:
            continue
        matched, _ = match_keywords(title, config)
        records.append({
            "record_type": "source", "origin": "legacy", "platform": "web", "content_type": "article",
            "url": link if link.startswith("http") else "",
            "url_kind": "permalink" if link.startswith("http") else "none",
            "container_name": outlet, "author_name": outlet, "author_url": None, "author_kind": "page",
            "posted_at_raw": "", "posted_at": None, "posted_at_precision": "unknown", "text": title,
            "attachments": [], "metrics": {}, "keywords_matched": matched,
            "entities_mentioned": [pid for pid, g in party_groups.items() if g and g in matched],
            "topics": [], "tone": None, "claim_type": None, "importance": "trung_binh", "importance_reason": None,
            "notes": "Chỉ có tiêu đề — nội dung bài sẽ được thu thập ở đợt Web",
            "legacy_ref": f"{SHEET_PRESS}!A{row}", "captured_at": LEGACY_AT,
        })
    return records


def _tally(results: list[dict]) -> dict:
    tally = {"added": 0, "duplicate": 0, "invalid": 0, "errors": []}
    for r in results:
        tally[r["status"]] += 1
        if r["status"] == "invalid":
            tally["errors"].append(r)
    return tally


def import_file(project: Project, path: Path | str) -> dict:
    path = Path(path)
    legacy_dir = project.root / "legacy"
    legacy_dir.mkdir(exist_ok=True)
    if not (legacy_dir / path.name).exists():
        shutil.copy2(path, legacy_dir / path.name)
    wb = load_workbook(path, data_only=True)

    config = project.load_config()
    merge_parties(config, wb[SHEET_PARTIES])
    events, disclaimer = event_records(wb[SHEET_EVENTS])
    if disclaimer:
        config["disclaimer"] = disclaimer
    containers = container_records(wb[SHEET_POSTS], config)
    merge_containers(config, containers)
    project.save_config(config)

    container_results = add_records(project, containers, LEGACY_RUN, config)
    container_ids = {normalize_url(c["url"]): r.get("id") or r.get("existing_id")
                     for c, r in zip(containers, container_results)}
    posts = post_records(wb[SHEET_POSTS], config, container_ids)
    press = press_records(wb[SHEET_PRESS], config)
    return {
        "containers": _tally(container_results),
        "sources_facebook": _tally(add_records(project, posts, LEGACY_RUN, config)),
        "sources_web": _tally(add_records(project, press, LEGACY_RUN, config)),
        "events": _tally(add_records(project, events, LEGACY_RUN, config)),
    }


def main(argv: list[str] | None = None) -> int:
    setup_stdout()
    parser = argparse.ArgumentParser(description="Nhập file Excel tổng hợp cũ vào kho")
    parser.add_argument("--project", help="thư mục dự án (mặc định: thư mục chứa .claude)")
    parser.add_argument("--file", required=True, help="đường dẫn file Excel cũ")
    args = parser.parse_args(argv)
    project = Project(args.project) if args.project else Project()
    print(json.dumps(import_file(project, args.file), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
